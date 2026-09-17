#!/usr/bin/env python3
"""
Genere le support PowerPoint de la session de sensibilisation.

    python3 scripts/support-powerpoint/generer_support.py

Le fichier produit est autonome : aucune connexion, aucun smartphone, aucune
installation n'est necessaire le jour de la session. Les notes du formateur
sont integrees a chaque diapositive et apparaissent en mode presentateur.
"""

import sys
from datetime import date
from pathlib import Path

from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

RACINE = Path(__file__).resolve().parent
sys.path.insert(0, str(RACINE))

from charte_cern import (  # noqa: E402
    ANTHRACITE,
    BLANC,
    BLEU_CERN,
    BLEU_NUIT,
    BLEU_PALE,
    COULEUR_ELEMENT,
    CYAN,
    GRIS_CLAIR,
    GRIS_MOYEN,
    GRIS_PALE,
    HAUT_CONTENU,
    HAUTEUR,
    LARGEUR,
    LARGEUR_UTILE,
    MARGE,
    ORANGE,
    POLICE,
    VIOLET,
    ajouter_bandeau,
    ajouter_bloc_texte,
    ajouter_encadre,
    ajouter_notes,
    ajouter_pied_de_page,
    ajouter_rectangle,
    ajouter_titre,
    estimer_hauteur_texte,
)
from contenu_elements import DIAPOS_ELEMENTS  # noqa: E402
from contenu_ouverture import DIAPOS_OUVERTURE  # noqa: E402
from contenu_synthese import DIAPOS_SYNTHESE  # noqa: E402
from visuels_verre import generer_visuels  # noqa: E402

MENTION_PIED = "LeVerre Labs  |  Sensibilisation au risque lié à l'activité physique"

COULEURS_NOMMEES = {
    "bleu": BLEU_CERN,
    "bleu_nuit": BLEU_NUIT,
    "cyan": CYAN,
    "violet": VIOLET,
    "orange": ORANGE,
    "gris": GRIS_MOYEN,
    **COULEUR_ELEMENT,
}


def _couleur(nom, defaut=BLEU_CERN):
    return COULEURS_NOMMEES.get(nom, defaut) if nom else defaut


def _lignes(entrees, taille=Pt(17), couleur=ANTHRACITE, puce=False):
    """Normalise une liste de chaines ou de dictionnaires en paragraphes."""
    resultat = []
    for entree in entrees:
        contenu = {"texte": entree} if isinstance(entree, str) else dict(entree)
        est_grande = contenu.pop("taille_grande", False)
        contenu.setdefault("taille", Pt(30) if est_grande else taille)
        contenu.setdefault("couleur", couleur)
        if puce and contenu["texte"] and not est_grande:
            contenu.setdefault("puce", True)
        resultat.append(contenu)
    return resultat


def _ajouter_source(diapo, texte):
    ajouter_bloc_texte(
        diapo,
        [{"texte": f"Source : {texte}", "taille": Pt(10), "couleur": GRIS_MOYEN, "italique": True}],
        haut=HAUTEUR - Inches(1.0),
        hauteur=Inches(0.3),
    )


# --- Gabarits ---------------------------------------------------------------


def gabarit_couverture(diapo, donnees, visuels):
    ajouter_bandeau(diapo, est_couverture=True)
    ajouter_rectangle(diapo, 0, Inches(0.13), LARGEUR, Inches(3.75), BLEU_NUIT)

    if "verre-moitie" in visuels:
        diapo.shapes.add_picture(
            str(visuels["verre-moitie"]),
            LARGEUR - Inches(3.4),
            Inches(0.55),
            height=Inches(2.85),
        )

    largeur_cartouche = Inches(8.9)
    ajouter_rectangle(diapo, MARGE, Inches(3.1), largeur_cartouche, Inches(1.5), BLANC)
    ajouter_bloc_texte(
        diapo,
        [{"texte": donnees["titre"], "taille": Pt(40), "gras": True, "couleur": BLEU_CERN}],
        gauche=Emu(MARGE.emu + Inches(0.4).emu),
        haut=Inches(3.42),
        largeur=Emu(largeur_cartouche.emu - Inches(0.8).emu),
        hauteur=Inches(0.9),
    )

    ajouter_bloc_texte(
        diapo,
        [{"texte": donnees["sous_titre"], "taille": Pt(20), "couleur": ANTHRACITE}],
        haut=Inches(5.0),
        hauteur=Inches(0.6),
    )
    ajouter_bloc_texte(
        diapo,
        _lignes(donnees.get("mentions", []), taille=Pt(14), couleur=GRIS_MOYEN),
        haut=Inches(5.75),
        hauteur=Inches(1.0),
        interligne=Pt(6),
    )


def gabarit_section(diapo, donnees, visuels):
    ajouter_rectangle(diapo, 0, 0, LARGEUR, HAUTEUR, BLEU_NUIT)
    for index, couleur in enumerate((BLEU_CERN, ORANGE, VIOLET, CYAN)):
        largeur_bloc = Emu(int(LARGEUR.emu / 4))
        ajouter_rectangle(
            diapo, Emu(index * largeur_bloc.emu), 0, largeur_bloc, Inches(0.13), couleur
        )

    ajouter_bloc_texte(
        diapo,
        [{"texte": donnees["numero"].upper(), "taille": Pt(16), "gras": True, "couleur": CYAN}],
        haut=Inches(2.4),
        hauteur=Inches(0.4),
    )
    ajouter_bloc_texte(
        diapo,
        [{"texte": donnees["titre"], "taille": Pt(44), "gras": True, "couleur": BLANC}],
        haut=Inches(2.9),
        hauteur=Inches(1.0),
    )
    if donnees.get("sous_titre"):
        ajouter_bloc_texte(
            diapo,
            [{"texte": donnees["sous_titre"], "taille": Pt(19), "couleur": GRIS_CLAIR}],
            haut=Inches(4.0),
            hauteur=Inches(0.5),
        )
    if donnees.get("duree"):
        ajouter_rectangle(diapo, MARGE, Inches(4.9), Inches(1.9), Inches(0.5), CYAN)
        ajouter_bloc_texte(
            diapo,
            [{"texte": donnees["duree"], "taille": Pt(15), "gras": True, "couleur": BLEU_NUIT}],
            gauche=MARGE,
            haut=Inches(5.02),
            largeur=Inches(1.9),
            hauteur=Inches(0.4),
            aligner=PP_ALIGN.CENTER,
        )


def gabarit_question(diapo, donnees, visuels):
    couleur = _couleur(donnees.get("couleur"))
    ajouter_rectangle(diapo, MARGE, Inches(2.15), Inches(1.4), Inches(0.08), couleur)
    ajouter_bloc_texte(
        diapo,
        [{"texte": donnees["question"], "taille": Pt(34), "gras": True, "couleur": couleur}],
        haut=Inches(2.6),
        hauteur=Inches(1.6),
    )
    lignes = [
        {"texte": ligne, "taille": Pt(19), "couleur": ANTHRACITE}
        for ligne in donnees.get("precision", "").split("\n")
        if ligne
    ]
    ajouter_bloc_texte(diapo, lignes, haut=Inches(4.35), hauteur=Inches(2.0), interligne=Pt(12))


def gabarit_chiffre(diapo, donnees, visuels):
    ajouter_bloc_texte(
        diapo,
        [{"texte": donnees["valeur"], "taille": Pt(96), "gras": True, "couleur": BLEU_CERN}],
        haut=Inches(1.5),
        hauteur=Inches(1.6),
        aligner=PP_ALIGN.CENTER,
    )
    ajouter_bloc_texte(
        diapo,
        [{"texte": donnees["legende"], "taille": Pt(22), "couleur": ANTHRACITE}],
        haut=Inches(3.15),
        hauteur=Inches(0.6),
        aligner=PP_ALIGN.CENTER,
    )
    ajouter_encadre(
        diapo,
        MARGE,
        Inches(4.1),
        LARGEUR_UTILE,
        Inches(1.5),
        _lignes(donnees.get("details", []), taille=Pt(18)),
    )
    if donnees.get("source"):
        _ajouter_source(diapo, donnees["source"])


def gabarit_puces(diapo, donnees, visuels):
    ajouter_titre(diapo, donnees["titre"], aligner=PP_ALIGN.LEFT)
    haut = HAUT_CONTENU
    if donnees.get("sous_titre"):
        ajouter_bloc_texte(
            diapo,
            [{"texte": donnees["sous_titre"], "taille": Pt(17), "italique": True, "couleur": GRIS_MOYEN}],
            haut=Inches(1.55),
            hauteur=Inches(0.4),
        )
        haut = Inches(2.15)

    ajouter_bloc_texte(
        diapo,
        _lignes(donnees["puces"], puce=True),
        haut=haut,
        hauteur=Inches(2.6),
        interligne=Pt(16),
    )
    if donnees.get("encadre"):
        lignes = _lignes(donnees["encadre"], taille=Pt(17))
        hauteur = Emu(
            estimer_hauteur_texte(
                lignes, Emu(LARGEUR_UTILE.emu - Inches(0.56).emu), Pt(17)
            ).emu
            + Inches(0.56).emu
        )
        bas = Inches(6.25) if donnees.get("source") else Inches(6.55)
        ajouter_encadre(
            diapo,
            MARGE,
            Emu(bas.emu - hauteur.emu),
            LARGEUR_UTILE,
            hauteur,
            lignes,
        )
    if donnees.get("source"):
        _ajouter_source(diapo, donnees["source"])


def _carte_colonne(diapo, gauche, haut, largeur, hauteur, titre, lignes, couleur):
    """Carte a en-tete colore. `hauteur=None` l'ajuste sur son contenu."""
    largeur_texte = Emu(int(largeur) - Inches(0.6).emu)
    if hauteur is None:
        hauteur = Emu(
            Inches(0.95).emu
            + estimer_hauteur_texte(lignes, largeur_texte, Pt(16), Pt(12)).emu
            + Inches(0.25).emu
        )

    ajouter_rectangle(diapo, gauche, haut, largeur, hauteur, GRIS_PALE)
    ajouter_rectangle(diapo, gauche, haut, largeur, Inches(0.62), couleur)
    ajouter_bloc_texte(
        diapo,
        [{"texte": titre, "taille": Pt(18), "gras": True, "couleur": BLANC}],
        gauche=Emu(int(gauche) + Inches(0.3).emu),
        haut=Emu(int(haut) + Inches(0.15).emu),
        largeur=largeur_texte,
        hauteur=Inches(0.4),
    )
    ajouter_bloc_texte(
        diapo,
        lignes,
        gauche=Emu(int(gauche) + Inches(0.3).emu),
        haut=Emu(int(haut) + Inches(0.95).emu),
        largeur=largeur_texte,
        hauteur=Emu(int(hauteur) - Inches(1.2).emu),
        interligne=Pt(12),
    )
    return Emu(int(hauteur))


def _ajuster_colonnes(donnees, largeur_texte, hauteur_maximale):
    """
    Choisit la plus grande taille de police qui tient dans la hauteur donnee.

    Sans cet ajustement, une colonne un peu trop fournie pousse le texte hors
    de sa carte ou par-dessus le pied de page.
    """
    for points in (16, 15, 14, 13):
        taille = Pt(points)
        gauche = _lignes(donnees["colonne_gauche"]["lignes"], taille=taille, puce=True)
        droite = _lignes(donnees["colonne_droite"]["lignes"], taille=taille, puce=True)
        hauteur = Emu(
            Inches(1.2).emu
            + max(
                estimer_hauteur_texte(gauche, largeur_texte, taille, Pt(12)).emu,
                estimer_hauteur_texte(droite, largeur_texte, taille, Pt(12)).emu,
            )
        )
        if hauteur.emu <= hauteur_maximale.emu or points == 13:
            return taille, hauteur, gauche, droite


def gabarit_colonnes(diapo, donnees, visuels):
    ajouter_titre(diapo, donnees["titre"])
    couleur = _couleur(donnees.get("couleur"))
    largeur_carte = Emu(int((LARGEUR_UTILE.emu - Inches(0.5).emu) / 2))
    haut = Inches(2.0)

    largeur_texte = Emu(largeur_carte.emu - Inches(0.6).emu)
    hauteur_maximale = Inches(3.5) if donnees.get("bas_de_page") else Inches(4.1)
    taille, hauteur, lignes_gauche, lignes_droite = _ajuster_colonnes(
        donnees, largeur_texte, hauteur_maximale
    )

    _carte_colonne(
        diapo,
        MARGE,
        haut,
        largeur_carte,
        hauteur,
        donnees["colonne_gauche"]["titre"],
        lignes_gauche,
        couleur,
    )
    _carte_colonne(
        diapo,
        Emu(MARGE.emu + largeur_carte.emu + Inches(0.5).emu),
        haut,
        largeur_carte,
        hauteur,
        donnees["colonne_droite"]["titre"],
        lignes_droite,
        BLEU_NUIT if couleur == BLEU_CERN else couleur,
    )

    if donnees.get("bas_de_page"):
        ajouter_bloc_texte(
            diapo,
            [{"texte": donnees["bas_de_page"], "taille": Pt(17), "gras": True, "couleur": BLEU_CERN}],
            haut=Emu(haut.emu + hauteur.emu + Inches(0.25).emu),
            hauteur=Inches(0.6),
            aligner=PP_ALIGN.CENTER,
        )
    if donnees.get("source"):
        _ajouter_source(diapo, donnees["source"])


def gabarit_visuel(diapo, donnees, visuels):
    ajouter_titre(diapo, donnees["titre"], aligner=PP_ALIGN.LEFT)
    diapo.shapes.add_picture(
        str(visuels[donnees["image"]]), LARGEUR - Inches(3.9), Inches(1.9), height=Inches(4.2)
    )
    ajouter_bloc_texte(
        diapo,
        _lignes(donnees["lignes"], taille=Pt(19)),
        haut=Inches(2.4),
        largeur=Inches(8.2),
        hauteur=Inches(3.4),
        interligne=Pt(14),
    )


def gabarit_modele(diapo, donnees, visuels):
    ajouter_titre(diapo, donnees["titre"])
    haut = Inches(1.85)
    for element in donnees["elements"]:
        couleur = COULEUR_ELEMENT[element["cle"]]
        ajouter_rectangle(diapo, MARGE, haut, Inches(0.16), Inches(0.5), couleur)
        ajouter_bloc_texte(
            diapo,
            [{"texte": element["nom"], "taille": Pt(17), "gras": True, "couleur": couleur}],
            gauche=Emu(MARGE.emu + Inches(0.35).emu),
            haut=Emu(haut.emu + Inches(0.05).emu),
            largeur=Inches(2.2),
            hauteur=Inches(0.4),
        )
        ajouter_bloc_texte(
            diapo,
            [{"texte": element["role"], "taille": Pt(16), "couleur": ANTHRACITE}],
            gauche=Emu(MARGE.emu + Inches(2.7).emu),
            haut=Emu(haut.emu + Inches(0.05).emu),
            largeur=Inches(8.7),
            hauteur=Inches(0.5),
        )
        haut = Emu(haut.emu + Inches(0.66).emu)

    ajouter_encadre(
        diapo,
        MARGE,
        Inches(5.35),
        LARGEUR_UTILE,
        Inches(1.15),
        _lignes(donnees["encadre"], taille=Pt(16)),
    )


def gabarit_element(diapo, donnees, visuels):
    couleur = COULEUR_ELEMENT[donnees["cle"]]
    ajouter_titre(diapo, donnees["nom"], couleur=couleur, aligner=PP_ALIGN.LEFT)
    ajouter_bloc_texte(
        diapo,
        [{"texte": donnees["sous_titre"], "taille": Pt(17), "italique": True, "couleur": GRIS_MOYEN}],
        haut=Inches(1.5),
        hauteur=Inches(0.4),
    )

    largeur_texte = Inches(8.5)
    haut_definition = Inches(2.05)
    hauteur_definition = ajouter_encadre(
        diapo,
        MARGE,
        haut_definition,
        largeur_texte,
        None,
        [{"texte": donnees["definition"], "taille": Pt(16), "couleur": ANTHRACITE}],
        fond=BLEU_PALE,
        barre=couleur,
    )
    ajouter_bloc_texte(
        diapo,
        _lignes(donnees["composantes"], taille=Pt(16), puce=True),
        haut=Emu(haut_definition.emu + hauteur_definition.emu + Inches(0.35).emu),
        largeur=largeur_texte,
        hauteur=Inches(2.6),
        interligne=Pt(11),
    )
    if donnees.get("image") in visuels:
        diapo.shapes.add_picture(
            str(visuels[donnees["image"]]), LARGEUR - Inches(3.3), Inches(2.1), height=Inches(3.9)
        )


def gabarit_notation(diapo, donnees, visuels):
    couleur = COULEUR_ELEMENT[donnees["cle"]]
    ajouter_titre(diapo, donnees["titre"], couleur=couleur, aligner=PP_ALIGN.LEFT)
    ajouter_bloc_texte(
        diapo,
        [{"texte": donnees["intitule"], "taille": Pt(17), "italique": True, "couleur": GRIS_MOYEN}],
        haut=Inches(1.5),
        hauteur=Inches(0.4),
    )

    haut = Inches(2.1)
    hauteur_ligne = Inches(0.58)
    for bareme in donnees["baremes"]:
        ajouter_rectangle(diapo, MARGE, haut, Inches(1.25), Inches(0.46), couleur)
        ajouter_bloc_texte(
            diapo,
            [{"texte": bareme["note"], "taille": Pt(16), "gras": True, "couleur": BLANC}],
            gauche=MARGE,
            haut=Emu(haut.emu + Inches(0.08).emu),
            largeur=Inches(1.25),
            hauteur=Inches(0.35),
            aligner=PP_ALIGN.CENTER,
        )
        ajouter_bloc_texte(
            diapo,
            [{"texte": bareme["libelle"], "taille": Pt(16), "couleur": ANTHRACITE}],
            gauche=Emu(MARGE.emu + Inches(1.6).emu),
            haut=Emu(haut.emu + Inches(0.08).emu),
            largeur=Inches(9.8),
            hauteur=Inches(0.45),
        )
        haut = Emu(haut.emu + hauteur_ligne.emu)

    haut_resultat = Emu(haut.emu + Inches(0.2).emu)
    ajouter_encadre(
        diapo,
        MARGE,
        haut_resultat,
        LARGEUR_UTILE,
        Inches(0.75),
        [{"texte": donnees["resultat"], "taille": Pt(18), "gras": True, "couleur": couleur}],
        fond=GRIS_PALE,
        barre=couleur,
    )
    if donnees.get("encadre"):
        ajouter_encadre(
            diapo,
            MARGE,
            Emu(haut_resultat.emu + Inches(0.95).emu),
            LARGEUR_UTILE,
            Inches(1.0),
            _lignes(donnees["encadre"], taille=Pt(15)),
        )


def gabarit_consigne(diapo, donnees, visuels):
    ajouter_titre(diapo, donnees["titre"])
    haut_consigne = Inches(1.75)
    hauteur_consigne = ajouter_encadre(
        diapo,
        MARGE,
        haut_consigne,
        LARGEUR_UTILE,
        None,
        _lignes(donnees["consigne"], taille=Pt(18)),
        taille=Pt(18),
    )
    haut_suite = Emu(haut_consigne.emu + hauteur_consigne.emu + Inches(0.45).emu)

    if donnees.get("colonnes_exemple"):
        largeur_carte = Emu(int((LARGEUR_UTILE.emu - Inches(0.5).emu) / 2))
        largeur_texte = Emu(largeur_carte.emu - Inches(0.6).emu)
        blocs = [donnees["colonnes_exemple"][cle] for cle in ("gauche", "droite")]
        lignes_par_bloc = [
            _lignes(bloc["lignes"], taille=Pt(15), puce=True, couleur=ANTHRACITE) for bloc in blocs
        ]
        hauteur_carte = Emu(
            Inches(1.2).emu
            + max(
                estimer_hauteur_texte(lignes, largeur_texte, Pt(15), Pt(12)).emu
                for lignes in lignes_par_bloc
            )
        )
        for index, (bloc, lignes) in enumerate(zip(blocs, lignes_par_bloc)):
            gauche = MARGE if index == 0 else Emu(MARGE.emu + largeur_carte.emu + Inches(0.5).emu)
            _carte_colonne(
                diapo,
                gauche,
                haut_suite,
                largeur_carte,
                hauteur_carte,
                bloc["titre"],
                lignes,
                _couleur(bloc.get("couleur")),
            )
    elif donnees.get("liste_simple"):
        haut = haut_suite
        for libelle in donnees["liste_simple"]:
            ajouter_rectangle(diapo, MARGE, haut, Inches(7.4), Inches(0.42), GRIS_PALE)
            ajouter_bloc_texte(
                diapo,
                [{"texte": libelle, "taille": Pt(16), "couleur": ANTHRACITE}],
                gauche=Emu(MARGE.emu + Inches(0.25).emu),
                haut=Emu(haut.emu + Inches(0.07).emu),
                largeur=Inches(5.5),
                hauteur=Inches(0.35),
            )
            ajouter_rectangle(diapo, Emu(MARGE.emu + Inches(7.6).emu), haut, Inches(1.3), Inches(0.42), BLANC)
            forme = diapo.shapes[-1]
            forme.line.color.rgb = GRIS_CLAIR
            forme.line.width = Pt(1)
            haut = Emu(haut.emu + Inches(0.52).emu)


def gabarit_calcul(diapo, donnees, visuels):
    ajouter_titre(diapo, donnees["titre"])
    exemple = donnees["exemple"]
    valeurs = {
        "robinet": f"{exemple['robinet']} × {str(exemple['bulle']).replace('.', ',')} = "
        f"{str(exemple['robinet'] * exemple['bulle']).replace('.0', '').replace('.', ',')}",
        "orage": f"+ {exemple['orage']}",
        "paille": f"− {exemple['paille']}",
    }
    resultat = exemple["robinet"] * exemple["bulle"] + exemple["orage"] - exemple["paille"]
    valeurs[None] = f"{str(resultat).replace('.0', '').replace('.', ',')} par heure"

    largeur_etape = Emu(int((LARGEUR_UTILE.emu - 3 * Inches(0.3).emu) / 4))
    for index, etape in enumerate(donnees["etapes"]):
        gauche = Emu(MARGE.emu + index * (largeur_etape.emu + Inches(0.3).emu))
        couleur = COULEUR_ELEMENT.get(etape["cle"], BLEU_NUIT)
        ajouter_rectangle(diapo, gauche, Inches(1.95), largeur_etape, Inches(2.05), GRIS_PALE)
        ajouter_rectangle(diapo, gauche, Inches(1.95), largeur_etape, Inches(0.09), couleur)
        ajouter_bloc_texte(
            diapo,
            [
                {"texte": etape["libelle"], "taille": Pt(17), "gras": True, "couleur": couleur},
                {"texte": etape["detail"], "taille": Pt(13), "couleur": GRIS_MOYEN},
            ],
            gauche=Emu(gauche.emu + Inches(0.22).emu),
            haut=Inches(2.22),
            largeur=Emu(largeur_etape.emu - Inches(0.44).emu),
            hauteur=Inches(1.3),
            interligne=Pt(6),
        )
        ajouter_bloc_texte(
            diapo,
            [{"texte": valeurs[etape["cle"]], "taille": Pt(16), "gras": True, "couleur": ANTHRACITE}],
            gauche=Emu(gauche.emu + Inches(0.22).emu),
            haut=Inches(3.45),
            largeur=Emu(largeur_etape.emu - Inches(0.44).emu),
            hauteur=Inches(0.4),
        )

    ajouter_bloc_texte(
        diapo,
        [{"texte": "Exemple : Robinet 6 · Bulle × 1,5 · Orage + 2 · Paille − 3 · Verre 40", "taille": Pt(14), "italique": True, "couleur": GRIS_MOYEN}],
        haut=Inches(4.2),
        hauteur=Inches(0.4),
        aligner=PP_ALIGN.CENTER,
    )
    ajouter_encadre(
        diapo,
        MARGE,
        Inches(4.75),
        LARGEUR_UTILE,
        Inches(1.45),
        [
            {"texte": donnees["formule_finale"], "taille": Pt(19), "gras": True, "couleur": BLEU_CERN},
            {
                "texte": f"Dans l'exemple : {exemple['verre']} ÷ "
                f"{str(resultat).replace('.0', '').replace('.', ',')} = 5 heures avant débordement",
                "taille": Pt(17),
                "couleur": ANTHRACITE,
            },
        ],
    )


def gabarit_lecture(diapo, donnees, visuels):
    ajouter_titre(diapo, donnees["titre"])
    largeur_carte = Emu(int((LARGEUR_UTILE.emu - 2 * Inches(0.35).emu) / 3))
    largeur_texte = Emu(largeur_carte.emu - Inches(0.6).emu)
    haut = Inches(2.05)
    hauteur_carte = Emu(
        Inches(1.2).emu
        + max(
            estimer_hauteur_texte(
                [{"texte": niveau["texte"], "taille": Pt(15)}], largeur_texte, Pt(15), Pt(12)
            ).emu
            for niveau in donnees["niveaux"]
        )
    )
    for index, niveau in enumerate(donnees["niveaux"]):
        gauche = Emu(MARGE.emu + index * (largeur_carte.emu + Inches(0.35).emu))
        _carte_colonne(
            diapo,
            gauche,
            haut,
            largeur_carte,
            hauteur_carte,
            niveau["titre"],
            [{"texte": niveau["texte"], "taille": Pt(15), "couleur": ANTHRACITE}],
            _couleur(niveau["couleur"]),
        )
    ajouter_encadre(
        diapo,
        MARGE,
        Emu(haut.emu + hauteur_carte.emu + Inches(0.4).emu),
        LARGEUR_UTILE,
        None,
        _lignes(donnees["encadre"], taille=Pt(17)),
        taille=Pt(17),
    )


def gabarit_cloture(diapo, donnees, visuels):
    ajouter_rectangle(diapo, 0, 0, LARGEUR, HAUTEUR, BLEU_NUIT)
    for index, couleur in enumerate((BLEU_CERN, ORANGE, VIOLET, CYAN)):
        largeur_bloc = Emu(int(LARGEUR.emu / 4))
        ajouter_rectangle(diapo, Emu(index * largeur_bloc.emu), 0, largeur_bloc, Inches(0.13), couleur)

    ajouter_bloc_texte(
        diapo,
        _lignes(donnees["lignes"], taille=Pt(18), couleur=BLANC),
        haut=Inches(1.7),
        hauteur=Inches(3.4),
        interligne=Pt(18),
    )
    ajouter_bloc_texte(
        diapo,
        [{"texte": donnees["titre"], "taille": Pt(40), "gras": True, "couleur": CYAN}],
        haut=Inches(5.4),
        hauteur=Inches(1.0),
    )


HAUT_BLOCS_FICHE = Inches(1.7)
BAS_BLOCS_FICHE = Inches(6.8)
ESPACE_BLOCS_FICHE = Inches(0.2)
GOUTTIERE_FICHE = Inches(0.4)


def gabarit_fiche(diapo, donnees, visuels):
    ajouter_titre(diapo, donnees["titre"], aligner=PP_ALIGN.LEFT, taille=Pt(24))
    ajouter_bloc_texte(
        diapo,
        [{"texte": donnees["sous_titre"], "taille": Pt(13), "italique": True, "couleur": GRIS_MOYEN}],
        haut=Inches(1.32),
        hauteur=Inches(0.3),
    )

    nombre_colonnes = donnees.get("colonnes", 2)
    largeur_colonne = Emu(
        int((LARGEUR_UTILE.emu - (nombre_colonnes - 1) * GOUTTIERE_FICHE.emu) / nombre_colonnes)
    )
    facteur, hauteurs, coupes = _repartir_blocs_fiche(
        donnees["blocs"], largeur_colonne, nombre_colonnes
    )

    haut = [HAUT_BLOCS_FICHE] * nombre_colonnes
    for index, (bloc, hauteur) in enumerate(zip(donnees["blocs"], hauteurs)):
        colonne = sum(1 for coupe in coupes if index >= coupe)
        gauche = Emu(MARGE.emu + colonne * (largeur_colonne.emu + GOUTTIERE_FICHE.emu))
        _dessiner_bloc_fiche(diapo, bloc, gauche, haut[colonne], largeur_colonne, hauteur, facteur)
        haut[colonne] = Emu(haut[colonne].emu + hauteur.emu + ESPACE_BLOCS_FICHE.emu)


def _repartir_blocs_fiche(blocs, largeur, nombre_colonnes):
    """
    Cherche la repartition en colonnes qui tient dans la page.

    On essaie d'abord a taille normale, puis on resserre progressivement les
    interlignes : mieux vaut une fiche un peu dense qu'un bloc hors de la page.
    """
    from itertools import combinations

    hauteur_disponible = Emu(BAS_BLOCS_FICHE - HAUT_BLOCS_FICHE)
    dernier = None

    for facteur in (1.0, 0.95, 0.9, 0.85, 0.8, 0.75):
        hauteurs = [_hauteur_bloc_fiche(bloc, largeur, facteur) for bloc in blocs]
        meilleur, ecart_minimal = None, None
        for coupes in combinations(range(1, len(blocs)), nombre_colonnes - 1):
            colonnes = _decouper(hauteurs, coupes)
            ecart = max(_hauteur_colonne(c) for c in colonnes) - min(
                _hauteur_colonne(c) for c in colonnes
            )
            if ecart_minimal is None or ecart < ecart_minimal:
                meilleur, ecart_minimal = coupes, ecart

        dernier = (facteur, hauteurs, meilleur)
        colonnes = _decouper(hauteurs, meilleur)
        if all(_hauteur_colonne(c) <= hauteur_disponible.emu for c in colonnes):
            return dernier

    return dernier


def _decouper(hauteurs, coupes):
    bornes = [0, *coupes, len(hauteurs)]
    return [hauteurs[bornes[i] : bornes[i + 1]] for i in range(len(bornes) - 1)]


def _hauteur_colonne(hauteurs):
    if not hauteurs:
        return 0
    return sum(h.emu for h in hauteurs) + (len(hauteurs) - 1) * ESPACE_BLOCS_FICHE.emu


def _hauteur_aide_fiche(bloc, largeur):
    """Les libelles d'aide font souvent deux lignes : on reserve la place."""
    if not bloc.get("aide"):
        return Inches(0.1)
    from charte_cern import estimer_nombre_lignes

    nombre = estimer_nombre_lignes(bloc["aide"], Emu(int(largeur) - Inches(0.36).emu), Pt(10))
    return Emu(Pt(nombre * 10 * 1.35).emu)


PAS_ITEM_FICHE = Inches(0.32)
PAS_LIGNE_FICHE = Inches(0.36)


def _hauteur_bloc_fiche(bloc, largeur, facteur=1.0):
    hauteur = Emu(Inches(0.44).emu + _hauteur_aide_fiche(bloc, largeur).emu + Inches(0.1).emu)
    hauteur = Emu(hauteur.emu + len(bloc.get("items", [])) * int(PAS_ITEM_FICHE.emu * facteur))
    hauteur = Emu(hauteur.emu + bloc.get("lignes_vides", 0) * int(PAS_LIGNE_FICHE.emu * facteur))
    if bloc.get("total"):
        hauteur = Emu(hauteur.emu + Inches(0.46).emu)
    return hauteur


def _case(diapo, gauche, haut, largeur=Inches(0.8), hauteur=Inches(0.28)):
    forme = ajouter_rectangle(diapo, gauche, haut, largeur, hauteur, BLANC)
    forme.line.color.rgb = GRIS_MOYEN
    forme.line.width = Pt(0.75)
    return forme


def _dessiner_bloc_fiche(diapo, bloc, gauche, haut, largeur, hauteur, facteur=1.0):
    couleur = _couleur(bloc.get("couleur"))
    ajouter_rectangle(diapo, gauche, haut, largeur, hauteur, GRIS_PALE)
    ajouter_rectangle(diapo, gauche, haut, largeur, Inches(0.36), couleur)
    ajouter_bloc_texte(
        diapo,
        [{"texte": bloc["titre"], "taille": Pt(13), "gras": True, "couleur": BLANC}],
        gauche=Emu(gauche.emu + Inches(0.18).emu),
        haut=Emu(haut.emu + Inches(0.06).emu),
        largeur=Emu(largeur.emu - Inches(0.36).emu),
        hauteur=Inches(0.3),
    )
    hauteur_aide = _hauteur_aide_fiche(bloc, largeur)
    if bloc.get("aide"):
        ajouter_bloc_texte(
            diapo,
            [{"texte": bloc["aide"], "taille": Pt(10), "italique": True, "couleur": GRIS_MOYEN}],
            gauche=Emu(int(gauche) + Inches(0.18).emu),
            haut=Emu(int(haut) + Inches(0.42).emu),
            largeur=Emu(int(largeur) - Inches(0.36).emu),
            hauteur=hauteur_aide,
            interligne=Pt(0),
        )

    curseur = Emu(int(haut) + Inches(0.44).emu + hauteur_aide.emu + Inches(0.1).emu)
    pas_item = Emu(int(PAS_ITEM_FICHE.emu * facteur))
    hauteur_case = Emu(int(pas_item.emu * 0.78))
    for item in bloc.get("items", []):
        ajouter_bloc_texte(
            diapo,
            [{"texte": item, "taille": Pt(11.5), "couleur": ANTHRACITE}],
            gauche=Emu(int(gauche) + Inches(0.18).emu),
            haut=curseur,
            largeur=Emu(int(largeur) - Inches(1.25).emu),
            hauteur=pas_item,
        )
        _case(
            diapo,
            Emu(int(gauche) + int(largeur) - Inches(0.98).emu),
            Emu(curseur.emu - Inches(0.02).emu),
            hauteur=hauteur_case,
        )
        curseur = Emu(curseur.emu + pas_item.emu)

    for _ in range(bloc.get("lignes_vides", 0)):
        ajouter_rectangle(
            diapo,
            Emu(int(gauche) + Inches(0.18).emu),
            Emu(curseur.emu + Inches(0.2).emu),
            Emu(int(largeur) - Inches(0.36).emu),
            Emu(Inches(0.012).emu),
            GRIS_CLAIR,
        )
        curseur = Emu(curseur.emu + int(PAS_LIGNE_FICHE.emu * facteur))

    if bloc.get("total"):
        ajouter_bloc_texte(
            diapo,
            [{"texte": bloc["total"], "taille": Pt(12), "gras": True, "couleur": couleur}],
            gauche=Emu(int(gauche) + Inches(0.18).emu),
            haut=Emu(curseur.emu + Inches(0.07).emu),
            largeur=Emu(int(largeur) - Inches(1.25).emu),
            hauteur=Inches(0.3),
        )
        _case(
            diapo,
            Emu(int(gauche) + int(largeur) - Inches(0.98).emu),
            Emu(curseur.emu + Inches(0.03).emu),
            hauteur=Inches(0.3),
        )


def gabarit_minutage(diapo, donnees, visuels):
    ajouter_titre(diapo, donnees["titre"], aligner=PP_ALIGN.LEFT)
    ajouter_bloc_texte(
        diapo,
        [{"texte": donnees["sous_titre"], "taille": Pt(14), "italique": True, "couleur": GRIS_MOYEN}],
        haut=Inches(1.5),
        hauteur=Inches(0.35),
    )
    haut = Inches(2.0)
    for index, (libelle, duree) in enumerate(donnees["etapes"]):
        fond = GRIS_PALE if index % 2 == 0 else BLANC
        ajouter_rectangle(diapo, MARGE, haut, LARGEUR_UTILE, Inches(0.44), fond)
        ajouter_bloc_texte(
            diapo,
            [{"texte": libelle, "taille": Pt(14), "couleur": ANTHRACITE}],
            gauche=Emu(MARGE.emu + Inches(0.25).emu),
            haut=Emu(haut.emu + Inches(0.09).emu),
            largeur=Inches(9.0),
            hauteur=Inches(0.3),
        )
        ajouter_bloc_texte(
            diapo,
            [{"texte": duree, "taille": Pt(14), "gras": True, "couleur": BLEU_CERN}],
            gauche=Emu(MARGE.emu + LARGEUR_UTILE.emu - Inches(1.4).emu),
            haut=Emu(haut.emu + Inches(0.09).emu),
            largeur=Inches(1.2),
            hauteur=Inches(0.3),
            aligner=PP_ALIGN.RIGHT,
        )
        haut = Emu(haut.emu + Inches(0.46).emu)


GABARITS = {
    "couverture": gabarit_couverture,
    "section": gabarit_section,
    "question": gabarit_question,
    "chiffre": gabarit_chiffre,
    "puces": gabarit_puces,
    "colonnes": gabarit_colonnes,
    "visuel": gabarit_visuel,
    "modele": gabarit_modele,
    "element": gabarit_element,
    "notation": gabarit_notation,
    "consigne": gabarit_consigne,
    "calcul": gabarit_calcul,
    "lecture": gabarit_lecture,
    "cloture": gabarit_cloture,
    "fiche": gabarit_fiche,
    "minutage": gabarit_minutage,
}

SANS_HABILLAGE = {"couverture", "section", "cloture"}


def construire(chemin_sortie, dossier_visuels):
    visuels = generer_visuels(dossier_visuels)

    presentation = Presentation()
    presentation.slide_width = LARGEUR
    presentation.slide_height = HAUTEUR
    mise_en_page_vide = presentation.slide_layouts[6]

    diapositives = DIAPOS_OUVERTURE + DIAPOS_ELEMENTS + DIAPOS_SYNTHESE
    for numero, donnees in enumerate(diapositives, start=1):
        diapo = presentation.slides.add_slide(mise_en_page_vide)
        gabarit = donnees["gabarit"]

        if gabarit not in SANS_HABILLAGE:
            ajouter_bandeau(diapo)
            ajouter_pied_de_page(diapo, numero, MENTION_PIED)

        GABARITS[gabarit](diapo, donnees, visuels)

        if donnees.get("notes"):
            ajouter_notes(diapo, "\n".join(l.strip() for l in donnees["notes"].splitlines()))

    Path(chemin_sortie).parent.mkdir(parents=True, exist_ok=True)
    presentation.save(chemin_sortie)
    return chemin_sortie, len(diapositives)


SORTIE_PAR_DEFAUT = "docs/supports/sensibilisation-tms-charte-cern.pptx"


if __name__ == "__main__":
    import tempfile

    sortie = sys.argv[1] if len(sys.argv) > 1 else SORTIE_PAR_DEFAUT
    # Les visuels sont regeneres a chaque execution : inutile de les versionner,
    # ils sont de toute facon embarques dans le fichier PowerPoint.
    with tempfile.TemporaryDirectory() as dossier_visuels:
        chemin, nombre = construire(sortie, sys.argv[2] if len(sys.argv) > 2 else dossier_visuels)
    print(f"{nombre} diapositives générées dans {chemin}")
