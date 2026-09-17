"""
Charte graphique CERN appliquee au support PowerPoint.

Sources publiques (aucun contenu sous authentification n'est utilise) :
- Palette officielle du template PowerPoint CERN, relevee sur l'image publiee
  par le Design and Visual Identity Service :
  design-guidelines.web.cern.ch/guidelines/ppt-colours
- Polices corporate (Optima / Helvetica Neue / Arial) :
  design-guidelines.web.cern.ch/guidelines/typography
  Arial est retenue : elle fait partie des polices corporate et elle est
  presente sur tous les postes, donc le rendu ne bouge pas d'un ordinateur
  a l'autre le jour de la session.

Le logo CERN et les visuels de l'Organisation ne sont volontairement pas
repris : la session est animee par un intervenant externe sous sa propre
marque. On reprend la charte, pas l'identite.
"""

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

# --- Palette officielle PowerPoint CERN -------------------------------------

BLEU_CERN = RGBColor(0x00, 0x33, 0xA0)
BLEU_NUIT = RGBColor(0x1C, 0x44, 0x6B)
CYAN = RGBColor(0x61, 0xC5, 0xD3)
VIOLET = RGBColor(0x6E, 0x24, 0x66)
ORANGE = RGBColor(0xE1, 0x5E, 0x32)
GRIS_CLAIR = RGBColor(0xBE, 0xBE, 0xC8)
GRIS_MOYEN = RGBColor(0x6D, 0x6E, 0x71)
ANTHRACITE = RGBColor(0x2F, 0x2F, 0x2F)
BLANC = RGBColor(0xFF, 0xFF, 0xFF)

# Fonds tres clairs derives de la palette, pour les encadres.
BLEU_PALE = RGBColor(0xE8, 0xEE, 0xF7)
GRIS_PALE = RGBColor(0xF2, 0xF2, 0xF4)

POLICE = "Arial"

# --- Format de diapositive ---------------------------------------------------

LARGEUR = Inches(13.333)
HAUTEUR = Inches(7.5)
MARGE = Inches(0.9)
LARGEUR_UTILE = Emu(LARGEUR - 2 * MARGE)

HAUTEUR_BANDEAU = Inches(0.13)
HAUT_TITRE = Inches(0.55)
HAUT_CONTENU = Inches(1.75)

# Couleurs metier des 5 elements, reexprimees dans la palette CERN pour que
# le support reste coherent avec la charte tout en gardant le modele lisible.
COULEUR_ELEMENT = {
    "verre": GRIS_MOYEN,
    "robinet": BLEU_CERN,
    "bulle": VIOLET,
    "orage": ORANGE,
    "paille": CYAN,
}


def _definir_texte(cadre, aligner=PP_ALIGN.LEFT):
    """Reglages communs a tous les cadres de texte."""
    cadre.word_wrap = True
    cadre.margin_left = 0
    cadre.margin_right = 0
    cadre.margin_top = 0
    cadre.margin_bottom = 0
    cadre.paragraphs[0].alignment = aligner
    return cadre


def _appliquer_puce(paragraphe, couleur, caractere="\u2022"):
    """Ajoute une vraie puce PowerPoint (et non un caractere tape a la main)."""
    pPr = paragraphe._p.get_or_add_pPr()
    pPr.set("marL", str(Emu(Inches(0.28)).emu))
    pPr.set("indent", str(-Emu(Inches(0.28)).emu))
    from lxml import etree

    ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
    couleur_xml = etree.SubElement(pPr, f"{{{ns}}}buClr")
    srgb = etree.SubElement(couleur_xml, f"{{{ns}}}srgbClr")
    srgb.set("val", str(couleur))
    police_xml = etree.SubElement(pPr, f"{{{ns}}}buFont")
    police_xml.set("typeface", "Arial")
    puce_xml = etree.SubElement(pPr, f"{{{ns}}}buChar")
    puce_xml.set("char", caractere)


def ajouter_rectangle(diapo, gauche, haut, largeur, hauteur, couleur, transparence=None):
    forme = diapo.shapes.add_shape(MSO_SHAPE.RECTANGLE, gauche, haut, largeur, hauteur)
    forme.fill.solid()
    forme.fill.fore_color.rgb = couleur
    forme.line.fill.background()
    forme.shadow.inherit = False
    if transparence is not None:
        _appliquer_transparence(forme, transparence)
    return forme


def _appliquer_transparence(forme, pourcentage):
    from lxml import etree

    ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
    remplissage = forme.fill._xPr.find(f"{{{ns}}}solidFill")
    couleur = remplissage.find(f"{{{ns}}}srgbClr")
    alpha = etree.SubElement(couleur, f"{{{ns}}}alpha")
    alpha.set("val", str(int((100 - pourcentage) * 1000)))


def ajouter_bandeau(diapo, est_couverture=False):
    """
    Bandeau superieur du template CERN : quatre blocs de couleur sur la
    diapositive de titre, bloc bleu nuit + aplat gris sur les diapositives
    de contenu.
    """
    if est_couverture:
        couleurs = [BLEU_NUIT, ORANGE, VIOLET, CYAN]
        largeur_bloc = Emu(int(LARGEUR.emu / 4))
        for index, couleur in enumerate(couleurs):
            ajouter_rectangle(
                diapo, Emu(index * largeur_bloc.emu), 0, largeur_bloc, HAUTEUR_BANDEAU, couleur
            )
        return

    largeur_bleue = Emu(int(LARGEUR.emu * 0.25))
    ajouter_rectangle(diapo, 0, 0, largeur_bleue, HAUTEUR_BANDEAU, BLEU_NUIT)
    ajouter_rectangle(
        diapo,
        largeur_bleue,
        0,
        Emu(LARGEUR.emu - largeur_bleue.emu),
        HAUTEUR_BANDEAU,
        GRIS_CLAIR,
    )


def ajouter_pied_de_page(diapo, numero, mention):
    """Pied de page discret, calque sur celui du template CERN."""
    cadre_gauche = diapo.shapes.add_textbox(
        MARGE, HAUTEUR - Inches(0.52), Inches(8), Inches(0.3)
    ).text_frame
    _definir_texte(cadre_gauche)
    passage = cadre_gauche.paragraphs[0].add_run()
    passage.text = mention
    passage.font.size = Pt(9)
    passage.font.name = POLICE
    passage.font.color.rgb = BLEU_CERN

    cadre_droit = diapo.shapes.add_textbox(
        LARGEUR - MARGE - Inches(2), HAUTEUR - Inches(0.52), Inches(2), Inches(0.3)
    ).text_frame
    _definir_texte(cadre_droit, PP_ALIGN.RIGHT)
    passage = cadre_droit.paragraphs[0].add_run()
    passage.text = str(numero)
    passage.font.size = Pt(9)
    passage.font.name = POLICE
    passage.font.color.rgb = BLEU_CERN


def ajouter_titre(diapo, texte, couleur=BLEU_CERN, taille=Pt(30), aligner=PP_ALIGN.CENTER):
    cadre = diapo.shapes.add_textbox(MARGE, HAUT_TITRE, LARGEUR_UTILE, Inches(0.95)).text_frame
    _definir_texte(cadre, aligner)
    cadre.vertical_anchor = MSO_ANCHOR.MIDDLE
    passage = cadre.paragraphs[0].add_run()
    passage.text = texte
    passage.font.size = taille
    passage.font.bold = True
    passage.font.name = POLICE
    passage.font.color.rgb = couleur
    return cadre


def ajouter_sous_titre(diapo, texte, couleur=GRIS_MOYEN, taille=Pt(15)):
    cadre = diapo.shapes.add_textbox(MARGE, Inches(1.42), LARGEUR_UTILE, Inches(0.4)).text_frame
    _definir_texte(cadre, PP_ALIGN.CENTER)
    passage = cadre.paragraphs[0].add_run()
    passage.text = texte
    passage.font.size = taille
    passage.font.italic = True
    passage.font.name = POLICE
    passage.font.color.rgb = couleur
    return cadre


def ajouter_bloc_texte(
    diapo,
    lignes,
    gauche=MARGE,
    haut=HAUT_CONTENU,
    largeur=None,
    hauteur=Inches(4.3),
    taille=Pt(17),
    couleur=ANTHRACITE,
    interligne=Pt(11),
    aligner=PP_ALIGN.LEFT,
):
    """
    `lignes` : liste de chaines, ou de dictionnaires acceptant les cles
    `texte`, `taille`, `gras`, `italique`, `couleur`, `puce`, `espace_avant`.
    """
    largeur = largeur if largeur is not None else LARGEUR_UTILE
    cadre = diapo.shapes.add_textbox(gauche, haut, largeur, hauteur).text_frame
    _definir_texte(cadre, aligner)

    for index, ligne in enumerate(lignes):
        contenu = {"texte": ligne} if isinstance(ligne, str) else dict(ligne)
        paragraphe = cadre.paragraphs[0] if index == 0 else cadre.add_paragraph()
        paragraphe.alignment = contenu.get("aligner", aligner)
        paragraphe.space_after = contenu.get("espace_apres", interligne)
        if contenu.get("espace_avant"):
            paragraphe.space_before = contenu["espace_avant"]

        couleur_ligne = contenu.get("couleur", couleur)
        if contenu.get("puce"):
            _appliquer_puce(paragraphe, contenu.get("couleur_puce", BLEU_CERN))

        passage = paragraphe.add_run()
        passage.text = contenu["texte"]
        passage.font.size = contenu.get("taille", taille)
        passage.font.bold = contenu.get("gras", False)
        passage.font.italic = contenu.get("italique", False)
        passage.font.name = POLICE
        passage.font.color.rgb = couleur_ligne

    return cadre


def estimer_nombre_lignes(texte, largeur_disponible, taille):
    """
    Estime le nombre de lignes qu'occupera un texte.

    PowerPoint calcule le retour a la ligne a l'ouverture du fichier, mais
    python-pptx ne le fait pas : sans cette estimation, les encadres sont
    dimensionnes a l'aveugle et le texte deborde. En Arial, la largeur moyenne
    d'un caractere vaut environ la moitie de la taille de police.
    """
    if not texte:
        return 1
    largeur_points = Emu(int(largeur_disponible)).inches * 72
    caracteres_par_ligne = max(1, int(largeur_points / (taille.pt * 0.52)))
    return max(1, -(-len(texte) // caracteres_par_ligne))


def estimer_hauteur_texte(lignes, largeur_disponible, taille_defaut, interligne=Pt(11)):
    """Hauteur totale d'un bloc de texte, interlignes compris."""
    total = 0
    for ligne in lignes:
        contenu = {"texte": ligne} if isinstance(ligne, str) else ligne
        taille = contenu.get("taille", taille_defaut)
        nombre = estimer_nombre_lignes(contenu["texte"], largeur_disponible, taille)
        total += nombre * taille.pt * 1.3 + contenu.get("espace_apres", interligne).pt
    return Emu(Pt(total).emu)


def ajouter_encadre(
    diapo,
    gauche,
    haut,
    largeur,
    hauteur,
    lignes,
    fond=BLEU_PALE,
    barre=BLEU_CERN,
    taille=Pt(16),
    padding=Inches(0.28),
):
    """
    Encadre a filet colore : sert aux consignes, definitions et exemples.

    `hauteur=None` dimensionne l'encadre sur son contenu. La hauteur retenue
    est renvoyee pour que l'appelant puisse enchainer l'element suivant.
    """
    largeur_texte = Emu(int(largeur) - 2 * padding.emu)
    if hauteur is None:
        hauteur = Emu(
            estimer_hauteur_texte(lignes, largeur_texte, taille).emu + 2 * padding.emu
        )

    ajouter_rectangle(diapo, gauche, haut, largeur, hauteur, fond)
    ajouter_rectangle(diapo, gauche, haut, Inches(0.07), hauteur, barre)
    ajouter_bloc_texte(
        diapo,
        lignes,
        gauche=Emu(int(gauche) + padding.emu),
        haut=Emu(int(haut) + padding.emu),
        largeur=largeur_texte,
        hauteur=Emu(int(hauteur) - 2 * padding.emu),
        taille=taille,
    )
    return Emu(int(hauteur))


def ajouter_notes(diapo, texte):
    """Notes du formateur, visibles en mode presentateur."""
    diapo.notes_slide.notes_text_frame.text = texte.strip()
